import csv

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2F7D32")
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Neutralise spreadsheet formulas in text (names come from devices and users): '=cmd' -> "'=cmd"."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def xlsx_response(filename, sheet_title, headers, rows, title=None, col_widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        start_row = 3
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            ws.cell(row=r, column=col, value=safe_cell(value))
    for col, h in enumerate(headers, 1):
        width = (col_widths or {}).get(col) or max(8, min(40, len(str(h)) + 4))
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    wb.save(resp)
    return resp


def csv_response(filename, headers, rows):
    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    w = csv.writer(resp)
    w.writerow(headers)
    w.writerows([safe_cell(v) for v in row] for row in rows)
    return resp
