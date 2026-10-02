"""Regression tests for the security review fixes."""
import datetime as dt
import os
import shutil
import tempfile

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from attendance.models import AttendanceLog, LeaveRequest, LeaveType
from core.excel import safe_cell
from employees.models import Employee

U = get_user_model()


class AccountTakeoverTests(TestCase):
    def test_staff_cannot_change_login_of_admin_account(self):
        boss = U.objects.create_superuser("boss", "", "Boss-pass-123")
        emp = Employee.objects.create(device_user_id="1", name="Boss", user=boss)
        hr = U.objects.create_user("hr", password="Hr-pass-1234", is_staff=True)
        self.client.force_login(hr)
        self.client.post(reverse("employees:login", args=[emp.pk]), {"username": "boss", "password": "Hacked-pass-99"})
        self.client.post(reverse("employees:login", args=[emp.pk]), {"disable": "1", "username": "boss"})
        boss.refresh_from_db()
        self.assertTrue(boss.check_password("Boss-pass-123"))
        self.assertTrue(boss.is_active)

    def test_weak_self_service_password_rejected(self):
        self.client.force_login(U.objects.create_superuser("a", "", "Admin-pass-1"))
        emp = Employee.objects.create(device_user_id="2", name="Asha")
        resp = self.client.post(reverse("employees:login", args=[emp.pk]), {"username": "asha", "password": "12345678"})
        self.assertEqual(resp.status_code, 200)
        emp.refresh_from_db()
        self.assertIsNone(emp.user)


class LoginProtectionTests(TestCase):
    def setUp(self):
        cache.clear()
        U.objects.create_superuser("owner", "", "Owner-pass-123")  # setup completed
        U.objects.create_user("emp", password="Right-pass-123")

    @override_settings(LOGIN_MAX_FAILURES=5)
    def test_lockout_after_failures(self):
        for i in range(5):
            self.client.post(reverse("login"), {"username": "emp", "password": f"wrong{i}"})
        resp = self.client.post(reverse("login"), {"username": "emp", "password": "Right-pass-123"})
        self.assertEqual(resp.status_code, 429)
        self.assertContains(resp, "Too many failed sign-in attempts", status_code=429)
        cache.clear()
        resp = self.client.post(reverse("login"), {"username": "emp", "password": "Right-pass-123"})
        self.assertEqual(resp.status_code, 302)

    def test_logout_needs_post(self):
        self.client.login(username="emp", password="Right-pass-123")
        self.client.get(reverse("logout_page"))
        self.assertIn("_auth_user_id", self.client.session)
        self.client.post(reverse("logout_page"))
        self.assertNotIn("_auth_user_id", self.client.session)


class RedirectAndInputTests(TestCase):
    def setUp(self):
        self.client.force_login(U.objects.create_superuser("a", "", "Admin-pass-1"))
        self.emp = Employee.objects.create(device_user_id="1", name="Asha")

    def test_open_redirects_blocked(self):
        for evil in ("https://evil.example/", "//evil.example/", "/\\evil.example/"):
            resp = self.client.post(reverse("attendance:manual_punch"), {
                "employee": self.emp.pk, "punch_time": "2026-10-01T10:00", "punch_state": 0, "note": "x", "next": evil})
            self.assertNotIn("evil", resp["Location"])
            log = AttendanceLog.objects.filter(source="manual").first()
            resp = self.client.post(reverse("attendance:delete_punch", args=[log.pk]), {"next": evil})
            self.assertNotIn("evil", resp["Location"])
        lt = LeaveType.objects.create(name="CL", code="CL")
        lr = LeaveRequest.objects.create(employee=self.emp, leave_type=lt, start_date=dt.date(2026, 11, 1), end_date=dt.date(2026, 11, 1))
        resp = self.client.post(reverse("attendance:leave_review", args=[lr.pk]), {"decision": "reject", "next": "https://evil.example/"})
        self.assertNotIn("evil", resp["Location"])

    def test_bad_parameters_do_not_crash(self):
        self.assertEqual(self.client.get(reverse("live_feed"), {"limit": "x"}).status_code, 200)
        self.assertEqual(self.client.get(reverse("live_feed"), {"limit": "-5"}).status_code, 200)
        self.assertEqual(self.client.get(reverse("attendance:logs"), {"device": "abc", "department": "x"}).status_code, 200)
        self.assertEqual(self.client.get(reverse("employees:list"), {"department": "x"}).status_code, 200)

    def test_device_address_must_be_ip(self):
        resp = self.client.post(reverse("devices:create"), {"name": "D", "ip_address": "evil.example", "port": 4370,
                                                            "comm_password": 0, "driver": "zk", "area": "A"})
        self.assertContains(resp, "Enter the device")

    def test_security_headers(self):
        resp = self.client.get(reverse("dashboard"))
        self.assertIn("frame-ancestors 'none'", resp["Content-Security-Policy"])
        self.assertEqual(resp["X-Frame-Options"], "DENY")


class SelfServiceTests(TestCase):
    def test_inactive_employee_loses_access(self):
        U.objects.create_superuser("owner", "", "Owner-pass-123")
        u = U.objects.create_user("e", password="Emp-pass-123")
        Employee.objects.create(device_user_id="5", name="Left", user=u, status="inactive")
        self.client.force_login(u)
        self.assertEqual(self.client.get(reverse("selfservice:home")).status_code, 403)


class SpreadsheetInjectionTests(TestCase):
    def test_formulas_neutralised(self):
        self.assertEqual(safe_cell('=HYPERLINK("http://x","y")'), '\'=HYPERLINK("http://x","y")')
        for v in ("+1", "-cmd", "@SUM(A1)"):
            self.assertTrue(safe_cell(v).startswith("'"))
        self.assertEqual(safe_cell("Asha"), "Asha")
        self.assertEqual(safe_cell(5), 5)
        self.client.force_login(U.objects.create_superuser("a", "", "Admin-pass-1"))
        Employee.objects.create(device_user_id="9", name='=cmd|" /C calc"!A0')
        resp = self.client.get(reverse("employees:list"), {"export": "xlsx"})
        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(resp.content)).active
        names = [c.value for c in ws["B"]]
        self.assertIn('\'=cmd|" /C calc"!A0', names)
        self.assertFalse(any(c.data_type == "f" for row in ws.iter_rows() for c in row))


class BackupHardeningTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.o = override_settings(BACKUP_WORK_DIR=self.tmp)
        self.o.enable()

    def tearDown(self):
        self.o.disable()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_public_status_is_generic(self):
        from backups import maintenance
        maintenance.enter(stage="download", label="secret-file.gz from Office NAS", pid=os.getpid())
        maintenance.write_status(message="pg_restore: connection to db host 10.0.0.5 role postgres failed")
        data = self.client.get("/backups/restore/status.json").json()
        self.assertNotIn("10.0.0.5", str(data))
        self.assertNotIn("Office NAS", str(data))
        maintenance.leave()

    def test_rclone_command_options_rejected(self):
        from backups.forms import DestinationForm
        base = {"name": "R", "kind": "rclone", "enabled": "on", "keep_last": 5, "rclone_remote": "box"}
        for bad in ("[box]\ntype = sftp\nssh = sh -c 'id'\n", "[box]\ntype = webdav\nbearer_token_command = id\n"):
            f = DestinationForm(data={**base, "rclone_config": bad})
            self.assertIn("rclone_config", f.errors)

    def test_backup_name_rejects_newline(self):
        from backups.transports import BACKUP_NAME_RE
        self.assertIsNone(BACKUP_NAME_RE.match("bioattend-backup-20260101-000000.gz\n"))


class HealthCheckTests(TestCase):
    def test_health_ok_without_host_or_login(self):
        resp = self.client.get("/healthz", HTTP_HOST="unknown-host.internal")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["status"], "ok")
        self.assertEqual(set(resp.json()), {"status", "database", "worker"})
